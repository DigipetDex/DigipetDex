import json
import os
import re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_sheet_name(name):
    # 이모지 및 특수문자 제거하여 엑셀/구글시트 완벽 호환
    cleaned = re.sub(r'[\U00010000-\U0010ffff]', '', name)
    cleaned = re.sub(r'[:\\/?*\[\]]', '_', cleaned).strip()
    if len(cleaned) > 30:
        cleaned = cleaned[:30]
    return cleaned or "Sheet"

def export_excel_by_dim():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    project_file = os.path.join(base_dir, "project_data.js")
    
    with open(project_file, "r", encoding="utf-8") as f:
        content = f.read().strip()
    
    json_str = content[content.index('{'):content.rindex('}')+1]
    data = json.loads(json_str)

    digimons = data.get("digimons", {})
    evolutions = data.get("evolutions", [])
    dims_list = data.get("dims", [])
    
    attr_map = {"vaccine": "백신", "data": "데이터", "virus": "바이러스", "free": "프리", "none": "-"}
    stage_order = {"디지타마": 0, "유년기1": 1, "유년기2": 2, "성장기": 3, "성숙기": 4, "완전체": 5, "궁극체": 6, "초궁극체": 7}

    # DiM별로 진화 데이터 그룹화
    dim_data = defaultdict(list)
    dim_digimons = defaultdict(set)

    for d_id, d in digimons.items():
        dim = d.get("dim") or "기타"
        dim_digimons[dim].add(d.get("name", d_id))

    for ev in evolutions:
        from_digi = digimons.get(ev.get("from"), {})
        to_digi = digimons.get(ev.get("to"), {})
        
        # DiM 결정 (to_digi 소속 우선)
        dim_name = to_digi.get("dim") or from_digi.get("dim") or "기타"
        
        req = ev.get("req") or to_digi.get("req") or {}
        time_val = req.get("time") or to_digi.get("time") or "-"
        if to_digi.get("unknownTime"):
            time_val = "불명"
            
        win_ratio = req.get("winRatio")
        win_str = f"{win_ratio}%" if (win_ratio is not None and win_ratio != "") else "-"

        jogress_val = req.get("jogress") or to_digi.get("jogress") or "-"

        row = {
            "from": from_digi.get("name", "미상"),
            "to": to_digi.get("name", "미상"),
            "stage": to_digi.get("stage", "-"),
            "attr": attr_map.get(to_digi.get("attr", ""), to_digi.get("attr", "-")),
            "vital": req.get("vital", "-"),
            "pp": req.get("pp", "-"),
            "battles": req.get("battles", "-"),
            "winRatio": win_str,
            "time": time_val,
            "jogress": jogress_val
        }
        dim_data[dim_name].append(row)

    # 엑셀 워크북 생성
    wb = openpyxl.Workbook()
    
    # 폰트 및 스타일 정의
    header_fill = PatternFill(start_color="3B4252", end_color="3B4252", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    
    summary_header_fill = PatternFill(start_color="434C5E", end_color="434C5E", fill_type="solid")
    
    data_font = Font(name="맑은 고딕", size=10)
    bold_font = Font(name="맑은 고딕", size=10, bold=True)
    
    thin_border_side = Side(border_style="thin", color="D8DEE9")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. 요약 시트 (DiM_목록_요약)
    ws_summary = wb.active
    ws_summary.title = "DiM_목록_요약"
    ws_summary.views.sheetView[0].showGridLines = True
    
    summary_headers = ["번호", "DiM 카드명", "등록 디지몬 수", "진화 루트 수"]
    ws_summary.append(summary_headers)
    for col_idx in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.fill = summary_header_fill
        cell.font = header_font
        cell.alignment = align_center

    # 등록된 전체 DiM 목록 정렬 (dims_list 우선, 그 외 추가된 DiM)
    all_dims = list(dims_list)
    for d in dim_data.keys():
        if d not in all_dims:
            all_dims.append(d)

    summary_row_idx = 2
    used_sheet_names = {"DiM_목록_요약"}

    for idx, dim in enumerate(all_dims, 1):
        d_count = len(dim_digimons.get(dim, []))
        e_count = len(dim_data.get(dim, []))
        ws_summary.append([idx, dim, d_count, e_count])
        for c in range(1, 5):
            cell = ws_summary.cell(row=summary_row_idx, column=c)
            cell.font = data_font
            cell.border = border_all
            cell.alignment = align_center if c != 2 else align_left
        summary_row_idx += 1

    # 2. 각 DiM별 시트 생성
    headers = ["진화 전", "진화 후", "세대", "속성", "바이탈", "PP/트로피", "배틀수", "승률", "진화시간", "조그레스 / 특이조건"]
    
    for dim in all_dims:
        base_title = clean_sheet_name(dim)
        title = base_title
        counter = 1
        while title in used_sheet_names:
            title = f"{base_title[:27]}_{counter}"
            counter += 1
        used_sheet_names.add(title)

        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        # 타이틀 & 헤더 행
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        # 해당 DiM의 진화 데이터 정렬
        ev_list = dim_data.get(dim, [])
        ev_list.sort(key=lambda r: (stage_order.get(r["stage"], 9), r["to"], r["from"]))

        for r_idx, ev in enumerate(ev_list, start=2):
            row_vals = [
                ev["from"],
                ev["to"],
                ev["stage"],
                ev["attr"],
                ev["vital"],
                ev["pp"],
                ev["battles"],
                ev["winRatio"],
                ev["time"],
                ev["jogress"]
            ]
            ws.append(row_vals)
            
            # 셀 스타일링
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = bold_font if c_idx in [1, 2] else data_font
                cell.border = border_all
                if c_idx in [1, 2]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

        # 열 너비 자동 조정
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                # 한글은 대략 1.8~2글자 너비 차지
                byte_len = sum(2 if ord(ch) > 127 else 1 for ch in val_str)
                if byte_len > max_len:
                    max_len = byte_len
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    # 요약 시트 열 너비 조정
    for col in ws_summary.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            byte_len = sum(2 if ord(ch) > 127 else 1 for ch in val_str)
            if byte_len > max_len:
                max_len = byte_len
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output_path = os.path.join(base_dir, "디지몬_DiM별_진화조건.xlsx")
    wb.save(output_path)
    print(f"성공! 총 {len(all_dims)}개 DiM 시트가 '{output_path}' 파일로 생성되었습니다.")

if __name__ == "__main__":
    export_excel_by_dim()
