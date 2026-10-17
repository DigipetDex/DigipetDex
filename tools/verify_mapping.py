# -*- coding: utf-8 -*-
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('디지몬랭크.xlsx', data_only=True)
ws = wb['Sheet1']

excel_digis = {}
for r in range(4, ws.max_row + 1):
    name = ws.cell(row=r, column=2).value
    hp = ws.cell(row=r, column=3).value
    ap = ws.cell(row=r, column=4).value
    bp = ws.cell(row=r, column=5).value
    if name and hp is not None and ap is not None and bp is not None:
        excel_digis[str(name).strip()] = {'hp': int(hp), 'ap': int(ap), 'spd': int(bp)}

for r in range(4, ws.max_row + 1):
    name = ws.cell(row=r, column=21).value
    hp = ws.cell(row=r, column=23).value
    ap = ws.cell(row=r, column=24).value
    bp = ws.cell(row=r, column=25).value
    if name and hp is not None and ap is not None and bp is not None:
        name_clean = str(name).strip()
        if name_clean not in excel_digis:
            excel_digis[name_clean] = {'hp': int(hp), 'ap': int(ap), 'spd': int(bp)}

with open('project_data.js', 'r', encoding='utf-8') as f:
    text = f.read()
start = text.find('{')
end = text.rfind('}') + 1
project = json.loads(text[start:end])
digimons = project.get('digimons', {})

NAME_ALIAS_MAP = {
    '미라쥬버스트': '미라쥬가오가몬: 버스트 모드',
    '벨젭블래스트': '베르제브몬: 블래스트 모드',
    '알포스X': '알포스브이드라몬(X항체)',
    '오메가S': '오메가몬 alter-S',
    '오메가몬B': '오메가몬 alter-B',
    '오메가몬B ': '오메가몬 alter-B',
    '오메가몬X': '오메가몬(X항체)',
    '카드라X': '카오스드라몬(X항체)',
    '헤클캅': '헤라클레스캅테리몬',
    '황팔': '황제드라몬: 팔라딘 모드',
    '타리언트캅테리몬': '타이런트캅테리몬',
    '레피드몬(황금)': '래피드몬(아머체)',
    '로제몬버스트': '로제몬: 버스트 모드',
    '듀크몬X': '듀크몬(X항체)',
    '듀크몬크림슨': '듀크몬: 크림슨 모드',
    '매그너몬X': '매그너몬(X항체)',
    '베르제브몬(C)': '베르제브몬(2010)',
    '샤우트몬DX': '샤우트몬 DX',
    '샤우트몬X7슈페리올모드': '샤우트몬x7: 슈페리올 모드',
    '샤인그레이몬:루인모드': '샤인그레이몬: 루인모드',
    '샤인그레이몬:버스트모드': '샤인그레이몬: 버스트 모드',
    '샤크라몬:무녀모드': '샤크라몬: 무녀 모드',
    '샤크라몬X': '샤크라몬(X항체)',
    '알파몬:왕룡검': '알파몬: 왕룡검',
    '레이브몬:버스트모드': '레이브몬: 버스트 모드',
    '펜리루가몬:타케미카즈치': '펜리루가몬: 타케미 카즈치',
    '황제드라몬:드래곤모드': '황제드라몬: 드래곤 모드',
    '벨페몬': '벨페몬: 레이지 모드',
    '에이션트이리시몬': '에이션트이리스몬',
    '파라사이몬': '패러사이몬',
}

proj_by_name = {}
for did, d in digimons.items():
    dname = d.get('name', '').strip()
    proj_by_name.setdefault(dname, []).append(d)

will_update = []
already_filled = []
not_in_project = []

for ex_name, st in sorted(excel_digis.items()):
    target_name = NAME_ALIAS_MAP.get(ex_name, ex_name)
    if target_name in proj_by_name:
        objs = proj_by_name[target_name]
        has_stats = any(d.get('baseHp') or d.get('baseAp') or d.get('baseSpd') for d in objs)
        if not has_stats:
            will_update.append((target_name, ex_name, len(objs), st))
        else:
            cur_stat = (objs[0].get('baseHp'), objs[0].get('baseAp'), objs[0].get('baseSpd'))
            already_filled.append((target_name, ex_name, cur_stat, st))
    else:
        not_in_project.append((ex_name, st))

print(f"Total Excel: {len(excel_digis)}")
print(f"1. [업데이트 대상 (미기입)] 총 {len(will_update)}개 종:")
for tname, exname, count, st in will_update:
    diff_note = f" (엑셀명: {exname})" if tname != exname else ""
    print(f"  * {tname}{diff_note} [{count}마리] -> HP: {st['hp']}, AP: {st['ap']}, SPD: {st['spd']}")

print(f"\n2. [이미 기입됨 (건너뜀)] 총 {len(already_filled)}개 종")

print(f"\n3. [프로젝트에 없는 디지몬] 총 {len(not_in_project)}개:")
for exname, st in not_in_project:
    print(f"  ? {exname}: {st}")
