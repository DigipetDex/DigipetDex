# -*- coding: utf-8 -*-
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('디지몬랭크.xlsx', data_only=True)
ws = wb['Sheet1']

excel_digis = {}

# Row 3 헤더 출력
row3_headers = {c: ws.cell(row=3, column=c).value for c in range(1, 30)}
print("Row 3 headers:", {k: v for k, v in row3_headers.items() if v})

# 왼쪽 테이블: B:이름(col 2), C:체력(col 3), D:공격(col 4), E:속도(col 5)
for r in range(4, ws.max_row + 1):
    name = ws.cell(row=r, column=2).value
    hp = ws.cell(row=r, column=3).value
    ap = ws.cell(row=r, column=4).value
    bp = ws.cell(row=r, column=5).value
    if name and hp is not None and ap is not None and bp is not None:
        name_clean = str(name).strip()
        excel_digis[name_clean] = {'hp': int(hp), 'ap': int(ap), 'spd': int(bp), 'src': 'left_table'}

# 오른쪽 테이블 컬럼 찾기
name_col = None
hp_col = None
ap_col = None
bp_col = None
for c in range(18, 30):
    val = ws.cell(row=3, column=c).value
    if val == '이름':
        name_col = c
    elif val == '체력':
        hp_col = c
    elif val == '공격':
        ap_col = c
    elif val == '속도':
        bp_col = c

print(f"Right table cols: name={name_col}, hp={hp_col}, ap={ap_col}, bp={bp_col}")

if name_col and hp_col and ap_col and bp_col:
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        hp = ws.cell(row=r, column=hp_col).value
        ap = ws.cell(row=r, column=ap_col).value
        bp = ws.cell(row=r, column=bp_col).value
        if name and hp is not None and ap is not None and bp is not None:
            name_clean = str(name).strip()
            if name_clean not in excel_digis:
                excel_digis[name_clean] = {'hp': int(hp), 'ap': int(ap), 'spd': int(bp), 'src': 'right_table'}
            else:
                print(f"Duplicate in right table: {name_clean}")

print(f"\nTotal digimons extracted from Excel: {len(excel_digis)}")

# project_data.js 로드
with open('project_data.js', 'r', encoding='utf-8') as f:
    text = f.read()

start = text.find('{')
end = text.rfind('}') + 1
project = json.loads(text[start:end])
digimons = project.get('digimons', {})

# project 내 디지몬 이름 맵 구성: name -> list of digi_objs
proj_name_map = {}
for did, d in digimons.items():
    dname = d.get('name', '').strip()
    if dname not in proj_name_map:
        proj_name_map[dname] = []
    proj_name_map[dname].append(d)

print(f"Total digimons in project_data.js: {len(digimons)}, unique names: {len(proj_name_map)}")

# 비교 분석
exact_match = []
fuzzy_candidates = []
not_in_project = []

for ex_name, ex_stat in sorted(excel_digis.items()):
    if ex_name in proj_name_map:
        objs = proj_name_map[ex_name]
        # 현재 스탯 상태 확인
        has_stats = any(d.get('baseHp') or d.get('baseAp') or d.get('baseSpd') for d in objs)
        exact_match.append({
            'excel_name': ex_name,
            'proj_name': ex_name,
            'count': len(objs),
            'has_stats': has_stats,
            'cur_hp': objs[0].get('baseHp'),
            'cur_ap': objs[0].get('baseAp'),
            'cur_spd': objs[0].get('baseSpd'),
            'new_stat': ex_stat
        })
    else:
        # 혹시 유사한 이름이 있는지 찾기
        found = False
        for pname in proj_name_map:
            # 띄어쓰기 제거나 축약어 매칭 검사
            ex_compact = ex_name.replace(' ', '').replace(':', '').replace('-', '')
            p_compact = pname.replace(' ', '').replace(':', '').replace('-', '')
            if ex_compact == p_compact or ex_compact in p_compact or p_compact in ex_compact:
                objs = proj_name_map[pname]
                has_stats = any(d.get('baseHp') or d.get('baseAp') or d.get('baseSpd') for d in objs)
                fuzzy_candidates.append({
                    'excel_name': ex_name,
                    'proj_name': pname,
                    'count': len(objs),
                    'has_stats': has_stats,
                    'cur_hp': objs[0].get('baseHp'),
                    'cur_ap': objs[0].get('baseAp'),
                    'cur_spd': objs[0].get('baseSpd'),
                    'new_stat': ex_stat
                })
                found = True
                break
        if not found:
            not_in_project.append((ex_name, ex_stat))

print("\n=== EXACT MATCHES (정확히 일치하는 이름) ===")
print(f"Total: {len(exact_match)}")
missing_stats_count = sum(1 for m in exact_match if not m['has_stats'])
print(f"  -> 미기입 상태인 것: {missing_stats_count}개")
print(f"  -> 이미 기입되어 있는 것: {len(exact_match) - missing_stats_count}개")

print("\n[미기입 디지몬 (정확 일치)]")
for m in exact_match:
    if not m['has_stats']:
        print(f"  {m['excel_name']} ({m['count']}마리) -> HP: {m['new_stat']['hp']}, AP: {m['new_stat']['ap']}, SPD: {m['new_stat']['spd']}")

print("\n=== FUZZY MATCH CANDIDATES (표기가 약간 다른 후보) ===")
print(f"Total: {len(fuzzy_candidates)}")
for f in fuzzy_candidates:
    status = "미기입" if not f['has_stats'] else f"기존기입 (HP:{f['cur_hp']}, AP:{f['cur_ap']}, SPD:{f['cur_spd']})"
    print(f"  Excel: '{f['excel_name']}' -> Project: '{f['proj_name']}' [{status}] | 신규 스탯: HP={f['new_stat']['hp']}, AP={f['new_stat']['ap']}, SPD={f['new_stat']['spd']}")

print("\n=== NOT FOUND IN PROJECT (프로젝트에 없는 이름) ===")
for n, st in not_in_project:
    print(f"  {n}: {st}")
